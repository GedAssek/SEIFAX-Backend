import openpyxl

file_path = r'c:\Users\AYENA Gédéon\OneDrive\Desktop\SEIFAX\Frontend\assets\Calculateur_Notes.xlsx'

try:
    wb = openpyxl.load_workbook(file_path, data_only=False)
    ws = wb['Feuil1']
    
    print("Formulas in Row 7 (Circuits électriques):")
    # Columns: 
    # E: Matière
    # F: Int Coef, G: Int Note
    # H: Comp Coef, I: Comp Note
    # J: Ex Coef, K: Ex Note
    # L: TP Coef, M: TP Note
    # N: Coef Total
    # O: Note finale
    # P: Statut
    # Q: Statut Groupe
    # R: Liens des copies
    
    row = 7
    print(f"N{row} (Coef Total):", ws[f'N{row}'].value)
    print(f"O{row} (Note finale):", ws[f'O{row}'].value)
    print(f"P{row} (Statut):", ws[f'P{row}'].value)
    print(f"Q{row} (Statut Groupe):", ws[f'Q{row}'].value)
    
    # Check Global Average (around row 50 probably)
    for r in range(40, 55):
        if ws[f'E{r}'].value == 'MOYENNE GENERALE':
            print(f"\nMoyenne Générale found at row {r}")
            print(f"N{r} (Total Coefs):", ws[f'N{r}'].value)
            print(f"O{r} (Moyenne):", ws[f'O{r}'].value)
            print(f"P{r} (Statut Global):", ws[f'P{r}'].value)
            
except Exception as e:
    print(f"Error: {e}")
