"""Import the official student spreadsheet into MongoDB.

Usage: python import_authorized_students.py "C:\\path\\LISTE DE LA PROMO EAC SEI 2025.xlsx"
The source spreadsheet is intentionally not copied into the repository.
"""
import os
import sys
from datetime import datetime, timezone

from dotenv import load_dotenv
from openpyxl import load_workbook
from pymongo import MongoClient, UpdateOne

from utils.identity import student_registry_key

PROMOTION_PREFIX = "EAC SEI 2025"


def main(path: str) -> None:
    if not os.path.isfile(path):
        raise SystemExit("Fichier Excel introuvable.")
    load_dotenv()
    mongo_url = os.getenv("MONGODB_URL")
    db_name = os.getenv("DB_NAME", "lefaxeur_db")
    if not mongo_url:
        raise SystemExit("MONGODB_URL doit être défini dans .env.")

    sheet = load_workbook(path, read_only=True, data_only=True).active
    operations = []
    for row in sheet.iter_rows(min_row=8, values_only=True):
        _, nom, prenom, sexe, groupe = row[:5]
        if not all((nom, prenom, sexe, groupe)):
            continue
        promotion = f"{PROMOTION_PREFIX} {str(groupe).strip().upper()}"
        key = student_registry_key(str(nom), str(prenom), str(sexe), promotion)
        if key["sexe"] not in {"M", "F"}:
            raise SystemExit(f"Sexe invalide dans le fichier pour {nom} {prenom}.")
        operations.append(UpdateOne(
            key,
            {"$setOnInsert": {**key, "nom": str(nom).strip(), "prenom": str(prenom).strip(),
                               "promotion": promotion, "created_at": datetime.now(timezone.utc).isoformat()}},
            upsert=True,
        ))

    client = MongoClient(mongo_url, serverSelectionTimeoutMS=10000)
    collection = client[db_name].authorized_students
    collection.create_index([
        ("nom_normalized", 1), ("prenom_normalized", 1), ("sexe", 1), ("promotion_normalized", 1)
    ], unique=True)
    result = collection.bulk_write(operations, ordered=False)
    print(f"Registre traité : {len(operations)} étudiants; {result.upserted_count} ajoutés.")


if __name__ == "__main__":
    if len(sys.argv) != 2:
        raise SystemExit("Usage : python import_authorized_students.py CHEMIN_DU_FICHIER.xlsx")
    main(sys.argv[1])
